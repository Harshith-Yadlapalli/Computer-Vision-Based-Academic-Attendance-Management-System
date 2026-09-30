import os
import sqlite3
import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

import app as app_module
import database as database_module
from openpyxl import load_workbook


class FacultyAttendanceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_path = os.path.join(self.temp_dir.name, "faculty-test.db")
        self.original_db = app_module.DB
        app_module.DB = self.database_path
        connection = sqlite3.connect(self.database_path)
        connection.executescript(
            """
            CREATE TABLE users (
                user_id TEXT PRIMARY KEY, password TEXT, role TEXT, dob TEXT,
                branch TEXT, name TEXT, year TEXT, section TEXT,
                faculty_approved INTEGER NOT NULL DEFAULT 0, email TEXT, phone TEXT,
                faculty_status TEXT NOT NULL DEFAULT 'pending'
            );
            CREATE TABLE attendance (
                id INTEGER PRIMARY KEY, student_id TEXT, date TEXT,
                first_seen_time TEXT, present INTEGER, class TEXT
            );
            CREATE TABLE period_attendance (
                id INTEGER PRIMARY KEY, student_id TEXT, date TEXT, period INTEGER,
                status TEXT, value REAL, first_seen_time TEXT,
                UNIQUE(student_id, date, period)
            );
            CREATE TABLE IF NOT EXISTS cctv_attendance (
                id INTEGER PRIMARY KEY, student_id TEXT, date TEXT, period INTEGER,
                room TEXT, first_seen_time TEXT, UNIQUE(student_id, date, period, room)
            );
            CREATE TABLE IF NOT EXISTS cctv_completed_periods (
                date TEXT, period INTEGER, room TEXT, completed_at TEXT,
                PRIMARY KEY(date, period, room)
            );
            CREATE TABLE faculty_attendance_submissions (
                id INTEGER PRIMARY KEY, faculty_id TEXT, class_id INTEGER, student_id TEXT,
                date TEXT, period INTEGER, status TEXT, submitted_at TEXT
            );
            CREATE TABLE attendance_review_flags (
                id INTEGER PRIMARY KEY, student_id TEXT, date TEXT, period INTEGER,
                flag_type TEXT, details TEXT, created_at TEXT,
                UNIQUE(student_id, date, period, flag_type)
            );
            CREATE TABLE class_schedules (
                id INTEGER PRIMARY KEY AUTOINCREMENT, branch TEXT NOT NULL, year TEXT NOT NULL,
                section TEXT NOT NULL, subject TEXT NOT NULL, weekday INTEGER NOT NULL,
                period INTEGER NOT NULL, room TEXT NOT NULL, faculty_id TEXT NOT NULL,
                UNIQUE(branch, year, section, subject, weekday, period)
            );
            CREATE TABLE faculty_classes (
                id INTEGER PRIMARY KEY AUTOINCREMENT, faculty_id TEXT NOT NULL,
                branch TEXT NOT NULL, year TEXT NOT NULL, section TEXT NOT NULL,
                UNIQUE(faculty_id, branch, year, section)
            );
            CREATE TABLE announcements (
                id INTEGER PRIMARY KEY AUTOINCREMENT, faculty_id TEXT NOT NULL, branch TEXT NOT NULL,
                year TEXT NOT NULL, section TEXT NOT NULL, title TEXT NOT NULL, body TEXT NOT NULL,
                created_at TEXT NOT NULL, audience TEXT NOT NULL DEFAULT 'all'
            );
            CREATE TABLE assignments (
                id INTEGER PRIMARY KEY AUTOINCREMENT, faculty_id TEXT NOT NULL, branch TEXT NOT NULL,
                year TEXT NOT NULL, section TEXT NOT NULL, title TEXT NOT NULL, description TEXT NOT NULL,
                due_at TEXT NOT NULL, created_at TEXT NOT NULL
            );
            CREATE TABLE assignment_submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT, assignment_id INTEGER NOT NULL,
                student_id TEXT NOT NULL, original_filename TEXT NOT NULL,
                stored_filename TEXT NOT NULL, submitted_at TEXT NOT NULL,
                UNIQUE(assignment_id, student_id)
            );
            CREATE TABLE leave_requests (
                id INTEGER PRIMARY KEY AUTOINCREMENT, requester_id TEXT NOT NULL,
                requester_role TEXT NOT NULL, branch TEXT NOT NULL, start_date TEXT NOT NULL,
                end_date TEXT NOT NULL, reason TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending',
                decision_note TEXT, created_at TEXT NOT NULL, reviewed_by TEXT, reviewed_at TEXT
            );
            CREATE TABLE faculty_attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT, faculty_id TEXT NOT NULL, branch TEXT NOT NULL,
                attendance_date TEXT NOT NULL, status TEXT NOT NULL, note TEXT,
                marked_by TEXT NOT NULL, marked_at TEXT NOT NULL,
                UNIQUE(faculty_id, attendance_date)
            );
            INSERT INTO users (user_id, password, role, dob, branch, name, year, section, faculty_approved, email, phone) VALUES
                ('STU001', 'pw', 'student', NULL, 'CSE', 'Student One', '2', 'A', 0, NULL, NULL),
                ('STU002', 'pw', 'student', NULL, 'CSE', 'Student Two', '2', 'A', 0, NULL, NULL),
                ('STU003', 'pw', 'student', NULL, 'CSE', 'Other Year', '3', 'A', 0, NULL, NULL),
                ('STU004', 'pw', 'student', NULL, 'ECE', 'Other Branch', '2', 'A', 0, NULL, NULL),
                ('STU005', 'pw', 'student', NULL, 'CSE', 'Student Five', '3', 'B', 0, NULL, NULL),
                ('cshod', 'pw', 'hod', NULL, 'CSE', NULL, NULL, NULL, 1, NULL, NULL);
            """
        )
        connection.commit()
        connection.close()
        self.patches = [
            patch.object(app_module.dbmod, "now_ist", return_value=datetime(2026, 9, 30, 10, 0)),
            patch.object(app_module.dbmod, "is_working_day", return_value=False),
            patch.object(app_module.database, "init_today", return_value=None),
            patch.object(app_module.database, "init_today_periods", return_value=None),
            patch.object(app_module.database, "backfill_all_students", return_value=None),
            patch.object(app_module.database, "backfill_student_periods", return_value=None),
        ]
        for active_patch in self.patches:
            active_patch.start()
        self.client = app_module.app.test_client()

    def tearDown(self):
        for active_patch in reversed(self.patches):
            active_patch.stop()
        app_module.DB = self.original_db
        self.temp_dir.cleanup()

    def set_role(self, client, user_id, role):
        with client.session_transaction() as session:
            session["user_id"] = user_id
            session["role"] = role

    def add_approved_faculty(self):
        connection = sqlite3.connect(self.database_path)
        connection.execute(
            "INSERT INTO users (user_id, password, role, branch, name, year, section, faculty_approved, faculty_status) "
            "VALUES ('FAC100', 'FacultyPass9!', 'faculty', 'CSE', 'Faculty One', '2', 'A', 1, 'approved')"
        )
        connection.execute(
            "INSERT INTO faculty_classes (faculty_id, branch, year, section) "
            "VALUES ('FAC100', 'CSE', '2', 'A')"
        )
        connection.commit()
        connection.close()

    def test_database_migration_preserves_existing_faculty_class_once(self):
        original_db_path = database_module.DB_PATH
        migration_path = os.path.join(self.temp_dir.name, "migration.db")
        connection = sqlite3.connect(migration_path)
        connection.execute(
            "CREATE TABLE users (user_id TEXT PRIMARY KEY, password TEXT, role TEXT, dob TEXT, "
            "branch TEXT, name TEXT, year TEXT, section TEXT, faculty_approved INTEGER)"
        )
        connection.execute(
            "INSERT INTO users VALUES ('FAC100', 'pw', 'faculty', NULL, 'CSE', 'Faculty One', '2', 'A', 1)"
        )
        connection.execute(
            "CREATE TABLE announcements (id INTEGER PRIMARY KEY, faculty_id TEXT, branch TEXT, year TEXT, "
            "section TEXT, title TEXT, body TEXT, created_at TEXT)"
        )
        connection.execute(
            "INSERT INTO announcements VALUES (1, 'cshod', 'ALL', 'ALL', 'ALL', 'Old notice', 'Still visible', '2026-09-30')"
        )
        connection.execute(
            "CREATE TABLE cctv_attendance (id INTEGER PRIMARY KEY, student_id TEXT, date TEXT, "
            "period INTEGER, first_seen_time TEXT, UNIQUE(student_id, date, period))"
        )
        connection.execute(
            "INSERT INTO cctv_attendance VALUES (1, 'STU001', '2026-09-29', 3, '09:40:00')"
        )
        connection.execute(
            "CREATE TABLE cctv_completed_periods (date TEXT, period INTEGER, completed_at TEXT, PRIMARY KEY(date, period))"
        )
        connection.execute(
            "INSERT INTO cctv_completed_periods VALUES ('2026-09-29', 3, '2026-09-29T10:20:00')"
        )
        connection.commit()
        connection.close()
        database_module.DB_PATH = migration_path
        try:
            with patch.object(database_module, "sync_holidays_google", return_value=None):
                database_module.init_db()
                database_module.init_db()
            connection = sqlite3.connect(migration_path)
            migrated_classes = connection.execute(
                "SELECT branch, year, section FROM faculty_classes WHERE faculty_id='FAC100'"
            ).fetchall()
            migrated_announcement = connection.execute(
                "SELECT title, audience FROM announcements WHERE id=1"
            ).fetchone()
            migrated_status = connection.execute(
                "SELECT faculty_status FROM users WHERE user_id='FAC100'"
            ).fetchone()[0]
            migrated_sighting = connection.execute(
                "SELECT student_id, date, period, room, first_seen_time FROM cctv_attendance"
            ).fetchone()
            migrated_completion = connection.execute(
                "SELECT date, period, room, completed_at FROM cctv_completed_periods"
            ).fetchone()
            connection.close()
            self.assertEqual(migrated_classes, [("CSE", "2", "A")])
            self.assertEqual(migrated_announcement, ("Old notice", "all"))
            self.assertEqual(migrated_status, "approved")
            self.assertEqual(migrated_sighting, ("STU001", "2026-09-29", 3, "", "09:40:00"))
            self.assertEqual(
                migrated_completion,
                ("2026-09-29", 3, "", "2026-09-29T10:20:00"),
            )
        finally:
            database_module.DB_PATH = original_db_path

    def test_hod_can_review_pending_approved_and_rejected_faculty(self):
        signup_data = {
            "name": "Faculty Request",
            "password": "FacultyPass9!",
            "branch": "CSE",
            "year": "2",
            "section": "A",
        }
        self.client.post("/faculty/signup", data={**signup_data, "user_id": "FAC-APPROVE"})
        self.client.post("/faculty/signup", data={**signup_data, "user_id": "FAC-REJECT"})
        self.set_role(self.client, "cshod", "hod")
        pending_page = self.client.get("/hod")
        self.assertIn(b"FAC-APPROVE", pending_page.data)
        self.assertIn(b"FAC-REJECT", pending_page.data)

        approved = self.client.post("/hod/approve-faculty", data={"user_id": "FAC-APPROVE"})
        rejected = self.client.post("/hod/reject-faculty", data={"user_id": "FAC-REJECT"})
        self.assertEqual(approved.status_code, 302)
        self.assertEqual(rejected.status_code, 302)
        reviewed_page = self.client.get("/hod")
        self.assertIn(b"FAC-APPROVE", reviewed_page.data)
        self.assertIn(b"Approved Faculty", reviewed_page.data)
        self.assertIn(b"Rejected Faculty", reviewed_page.data)
        self.assertIn(b"FAC-REJECT", reviewed_page.data)
        self.assertIn(b"No pending faculty requests.", reviewed_page.data)
        rejected_login = self.client.post(
            "/",
            data={"user_id": "FAC-REJECT", "password": "FacultyPass9!", "role": "faculty"},
        )
        self.assertIn(b"signup request was rejected", rejected_login.data)

    def test_cctv_reconciliation_updates_only_after_period_completion(self):
        connection = sqlite3.connect(self.database_path)
        connection.executescript(
            """
            INSERT INTO period_attendance (student_id, date, period, status, value, first_seen_time)
            VALUES ('STU001', '2026-09-30', 3, 'present', 1.0, '09:30:00'),
                   ('STU002', '2026-09-30', 3, 'present', 1.0, '09:31:00');
            INSERT INTO cctv_attendance (student_id, date, period, room, first_seen_time)
            VALUES ('STU001', '2026-09-30', 3, 'C-305', '09:35:00');
            """
        )
        connection.commit()
        before_reconciliation = connection.execute(
            "SELECT student_id, status FROM period_attendance WHERE date='2026-09-30' ORDER BY student_id"
        ).fetchall()
        connection.close()
        self.assertEqual(before_reconciliation, [("STU001", "present"), ("STU002", "present")])

        original_db_path = database_module.DB_PATH
        database_module.DB_PATH = self.database_path
        try:
            discrepancies = database_module.reconcile_cctv_period(
                "2026-09-30", 3, "C-305", ["STU001", "STU002"]
            )
        finally:
            database_module.DB_PATH = original_db_path

        connection = sqlite3.connect(self.database_path)
        records = connection.execute(
            "SELECT student_id, status, value, first_seen_time FROM period_attendance "
            "WHERE date='2026-09-30' AND period=3 ORDER BY student_id"
        ).fetchall()
        completed = connection.execute(
            "SELECT date, period, room FROM cctv_completed_periods"
        ).fetchall()
        connection.close()
        self.assertEqual(discrepancies, ["STU002"])
        self.assertEqual(
            records,
            [("STU001", "present", 1.0, "09:35:00"), ("STU002", "absent", 0.0, None)],
        )
        self.assertEqual(completed, [("2026-09-30", 3, "C-305")])

    def test_completed_cctv_overrides_faculty_and_review_flags_are_visible(self):
        self.add_approved_faculty()
        connection = sqlite3.connect(self.database_path)
        connection.execute(
            "INSERT INTO cctv_completed_periods (date, period, room, completed_at) "
            "VALUES ('2026-09-29', 3, 'C-305', '2026-09-29T10:20:00')"
        )
        connection.execute(
            "INSERT INTO cctv_attendance (student_id, date, period, room, first_seen_time) "
            "VALUES ('STU001', '2026-09-29', 3, 'C-305', '09:40:00')"
        )
        connection.executemany(
            "INSERT INTO class_schedules (branch, year, section, subject, weekday, period, room, faculty_id) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            [
                ("CSE", "2", "A", "Data Structures", 1, 3, "C-305", "FAC100"),
                ("CSE", "3", "B", "Algorithms", 1, 3, "C-306", "FAC100"),
            ],
        )
        connection.commit()
        connection.close()
        self.set_role(self.client, "FAC100", "faculty")

        post_data = {
            "date": "2026-09-29",
            "period": "3",
            "attendance_STU001": "present",
            "attendance_STU002": "present",
        }
        first_post = self.client.post("/faculty", data=post_data)
        self.assertEqual(first_post.status_code, 302)
        repeated_post = self.client.post("/faculty", data=post_data)
        self.assertEqual(repeated_post.status_code, 302)

        connection = sqlite3.connect(self.database_path)
        records = connection.execute(
            "SELECT student_id, status, value FROM period_attendance "
            "WHERE date='2026-09-29' AND period=3 ORDER BY student_id"
        ).fetchall()
        flags = connection.execute(
            "SELECT student_id, flag_type FROM attendance_review_flags "
            "WHERE date='2026-09-29' AND period=3 ORDER BY student_id, flag_type"
        ).fetchall()
        connection.close()
        self.assertEqual(records, [("STU001", "present", 1.0), ("STU002", "absent", 0.0)])
        self.assertEqual(
            flags,
            [
                ("STU001", "duplicate_submission"),
                ("STU002", "cctv_mismatch"),
                ("STU002", "duplicate_submission"),
            ],
        )

        self.set_role(self.client, "cshod", "hod")
        review_page = self.client.get("/hod?attendance_date=2026-09-29")
        self.assertIn(b"Attendance review", review_page.data)
        self.assertIn(b"Possible proxy", review_page.data)
        self.assertIn(b"Faculty/CCTV mismatch", review_page.data)
        self.assertIn(b"Daily CCTV attendance audit", review_page.data)
        self.assertIn(b"Seen", review_page.data)
        self.assertIn(b"Not seen", review_page.data)
        self.assertIn(b"Not verified", review_page.data)

    def test_home_and_sign_in_links_work_for_every_role(self):
        expected_home = {"hod": "/hod", "faculty": "/faculty", "student": "/student"}
        for role, home_path in expected_home.items():
            client = app_module.app.test_client()
            self.set_role(client, f"{role}-user", role)
            home = client.get("/home")
            self.assertEqual(home.location, home_path)
            sign_in = client.get("/login")
            self.assertEqual(sign_in.location, "/")
            with client.session_transaction() as session:
                self.assertNotIn("user_id", session)

        self.set_role(self.client, "cshod", "hod")
        dashboard = self.client.get("/hod")
        self.assertIn(b'href="/home">Home</a>', dashboard.data)
        self.assertNotIn(b'href="/login">Sign in</a>', dashboard.data)

    def test_student_and_faculty_leave_requests_show_hod_decisions(self):
        self.add_approved_faculty()
        student_client = app_module.app.test_client()
        self.set_role(student_client, "STU001", "student")
        student_submit = student_client.post(
            "/leave-request",
            data={"start_date": "2026-10-01", "end_date": "2026-10-02", "reason": "Medical appointment"},
        )
        self.assertEqual(student_submit.location, "/student?leave=saved")
        connection = sqlite3.connect(self.database_path)
        student_request_id = connection.execute(
            "SELECT id FROM leave_requests WHERE requester_id='STU001'"
        ).fetchone()[0]
        connection.close()

        self.set_role(self.client, "cshod", "hod")
        hod_pending = self.client.get("/hod")
        self.assertIn(b"Student and Faculty Leave Requests", hod_pending.data)
        self.assertIn(b"Medical appointment", hod_pending.data)
        approved = self.client.post(
            f"/hod/leave-requests/{student_request_id}/decision",
            data={"decision": "approved", "decision_note": "Approved"},
        )
        self.assertEqual(approved.status_code, 302)
        student_page = student_client.get("/student")
        self.assertIn(b"Approved", student_page.data)
        self.assertIn(b"Approved", student_page.data)

        faculty_client = app_module.app.test_client()
        self.set_role(faculty_client, "FAC100", "faculty")
        faculty_submit = faculty_client.post(
            "/leave-request",
            data={"start_date": "2026-10-03", "end_date": "2026-10-03", "reason": "Personal leave"},
        )
        self.assertEqual(faculty_submit.location, "/faculty?leave=saved")
        connection = sqlite3.connect(self.database_path)
        faculty_request_id = connection.execute(
            "SELECT id FROM leave_requests WHERE requester_id='FAC100'"
        ).fetchone()[0]
        connection.close()
        hod_faculty_request = self.client.get("/hod")
        self.assertIn(b"Personal leave", hod_faculty_request.data)
        rejected = self.client.post(
            f"/hod/leave-requests/{faculty_request_id}/decision",
            data={"decision": "rejected", "decision_note": "Please contact your HOD"},
        )
        self.assertEqual(rejected.status_code, 302)
        faculty_page = faculty_client.get("/faculty")
        self.assertIn(b"Rejected", faculty_page.data)
        self.assertIn(b"Please contact your HOD", faculty_page.data)

    def test_hod_records_department_faculty_attendance_by_date(self):
        self.add_approved_faculty()
        self.set_role(self.client, "cshod", "hod")
        page = self.client.get("/hod/faculty-attendance?date=2026-09-30")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Faculty One", page.data)
        saved = self.client.post(
            "/hod/faculty-attendance",
            data={
                "faculty_id": "FAC100",
                "date": "2026-09-30",
                "status": "present",
                "note": "On campus",
            },
        )
        self.assertEqual(saved.location, "/hod/faculty-attendance?date=2026-09-30&saved=1")
        connection = sqlite3.connect(self.database_path)
        record = connection.execute(
            "SELECT faculty_id, branch, attendance_date, status, note, marked_by "
            "FROM faculty_attendance WHERE faculty_id='FAC100'"
        ).fetchone()
        connection.close()
        self.assertEqual(record, ("FAC100", "CSE", "2026-09-30", "present", "On campus", "cshod"))

    def test_hod_leave_and_faculty_attendance_are_department_scoped(self):
        self.add_approved_faculty()
        connection = sqlite3.connect(self.database_path)
        connection.execute(
            "INSERT INTO users (user_id, password, role, branch, name, year, section, faculty_approved, faculty_status) "
            "VALUES ('FAC-DS', 'pw', 'faculty', 'DS', 'Faculty DS', '4', 'B', 1, 'approved')"
        )
        connection.execute(
            "INSERT INTO leave_requests (requester_id, requester_role, branch, start_date, end_date, reason, created_at) "
            "VALUES ('FAC-DS', 'faculty', 'DS', '2026-10-01', '2026-10-01', 'DS only request', '2026-09-30')"
        )
        connection.commit()
        ds_request_id = connection.execute(
            "SELECT id FROM leave_requests WHERE reason='DS only request'"
        ).fetchone()[0]
        connection.close()

        self.set_role(self.client, "cshod", "hod")
        cse_page = self.client.get("/hod")
        self.assertNotIn(b"DS only request", cse_page.data)
        forbidden = self.client.post(
            f"/hod/leave-requests/{ds_request_id}/decision", data={"decision": "approved"}
        )
        self.assertEqual(forbidden.status_code, 403)
        cse_faculty_page = self.client.get("/hod/faculty-attendance")
        self.assertNotIn(b"Faculty DS", cse_faculty_page.data)

        ds_hod_client = app_module.app.test_client()
        self.set_role(ds_hod_client, "dshod", "hod")
        ds_page = ds_hod_client.get("/hod")
        self.assertIn(b"DS only request", ds_page.data)
        approved = ds_hod_client.post(
            f"/hod/leave-requests/{ds_request_id}/decision", data={"decision": "approved"}
        )
        self.assertEqual(approved.status_code, 302)
        ds_faculty_page = ds_hod_client.get("/hod/faculty-attendance")
        self.assertIn(b"Faculty DS", ds_faculty_page.data)

    def test_hod_signup_redirects_to_sign_in_with_confirmation(self):
        signup = self.client.post(
            "/hod/signup",
            data={"branch_code": "it", "password": "FacultyPass9!"},
        )
        self.assertEqual(signup.status_code, 302)
        self.assertEqual(signup.location, "/?signup=hod")
        sign_in = self.client.get(signup.location)
        self.assertIn(b"HOD account created", sign_in.data)

    def test_approved_faculty_marks_only_assigned_class_for_shared_dashboards(self):
        signup = self.client.post(
            "/faculty/signup",
            data={
                "user_id": "FAC100",
                "name": "Faculty CSE",
                "password": "FacultyPass9!",
                "branch": "CSE",
                "year": "2",
                "section": "A",
            },
        )
        self.assertEqual(signup.status_code, 302)
        self.assertEqual(signup.location, "/?signup=faculty")
        signup_page = self.client.get(signup.location)
        self.assertIn(b"Faculty signup submitted", signup_page.data)

        pending_login = self.client.post(
            "/",
            data={"user_id": "FAC100", "password": "FacultyPass9!", "role": "faculty"},
        )
        self.assertIn(b"waiting for HOD approval", pending_login.data)

        hod_client = app_module.app.test_client()
        self.set_role(hod_client, "cshod", "hod")
        pending_page = hod_client.get("/hod")
        self.assertIn(b"FAC100", pending_page.data)
        approved = hod_client.post("/hod/approve-faculty", data={"user_id": "FAC100"})
        self.assertEqual(approved.status_code, 302)

        faculty_login = self.client.post(
            "/",
            data={"user_id": "FAC100", "password": "FacultyPass9!", "role": "faculty"},
        )
        self.assertEqual(faculty_login.location, "/faculty")
        roster_page = self.client.get("/faculty")
        self.assertIn(b"STU001", roster_page.data)
        self.assertIn(b"STU002", roster_page.data)
        self.assertNotIn(b"STU003", roster_page.data)
        self.assertNotIn(b"STU004", roster_page.data)

        saved = self.client.post(
            "/faculty",
            data={
                "date": "2026-09-29",
                "period": "3",
                "attendance_STU001": "present",
                "attendance_STU002": "od",
                "attendance_STU004": "present",
            },
        )
        self.assertEqual(saved.status_code, 302)

        connection = sqlite3.connect(self.database_path)
        records = connection.execute(
            "SELECT student_id, status, value FROM period_attendance ORDER BY student_id"
        ).fetchall()
        connection.close()
        self.assertEqual(
            records,
            [("STU001", "present", 1.0), ("STU002", "present", 1.0)],
        )

        student_client = app_module.app.test_client()
        self.set_role(student_client, "STU001", "student")
        student_page = student_client.get("/student")
        self.assertEqual(student_page.status_code, 200)
        self.assertIn(b"2026-09-29", student_page.data)

        hod_page = hod_client.get("/hod")
        self.assertEqual(hod_page.status_code, 200)
        self.assertIn(b"STU001", hod_page.data)
        self.assertIn(b"STU002", hod_page.data)
        self.assertIn(b"Approved Faculty", hod_page.data)
        self.assertIn(b"Faculty CSE", hod_page.data)

    def test_faculty_profile_and_student_directory_show_class_data(self):
        self.add_approved_faculty()
        connection = sqlite3.connect(self.database_path)
        connection.executemany(
            "INSERT INTO period_attendance (student_id, date, period, status, value) VALUES (?, '2026-09-29', ?, ?, ?)",
            [("STU001", 1, "present", 1.0), ("STU001", 2, "absent", 0.0)],
        )
        connection.commit()
        connection.close()
        self.set_role(self.client, "FAC100", "faculty")

        directory = self.client.get("/faculty/students")
        self.assertEqual(directory.status_code, 200)
        self.assertIn(b"Student One", directory.data)
        self.assertIn(b"50.0%", directory.data)
        report = self.client.get("/faculty/export-excel")
        self.assertEqual(report.status_code, 200)
        workbook = load_workbook(BytesIO(report.data), read_only=True, data_only=True)
        report_rows = list(workbook["Class Attendance"].values)
        workbook.close()
        self.assertTrue(any(row[0] == "STU001" and row[5] == 50 for row in report_rows))
        profile = self.client.get("/faculty/profile")
        self.assertEqual(profile.status_code, 200)
        updated = self.client.post(
            "/faculty/profile",
            data={
                "name": "Faculty Updated",
                "email": "faculty@example.edu",
                "phone": "1234567890",
                "current_password": "FacultyPass9!",
                "new_password": "NewFacultyPass9!",
            },
        )
        self.assertEqual(updated.status_code, 302)
        connection = sqlite3.connect(self.database_path)
        profile_values = connection.execute(
            "SELECT name, email, phone, password FROM users WHERE user_id='FAC100'"
        ).fetchone()
        connection.close()
        self.assertEqual(profile_values, ("Faculty Updated", "faculty@example.edu", "1234567890", "NewFacultyPass9!"))

    def test_hod_publishes_schedule_for_approved_class(self):
        self.add_approved_faculty()
        self.set_role(self.client, "cshod", "hod")
        connection = sqlite3.connect(self.database_path)
        class_id = connection.execute(
            "SELECT id FROM faculty_classes WHERE faculty_id='FAC100' AND year='2' AND section='A'"
        ).fetchone()[0]
        connection.close()
        saved = self.client.post(
            "/hod/classes",
            data={
                "class_assignment_id": str(class_id),
                "subject": "Data Structures",
                "room": "C-204",
                "weekday": "0",
                "period": "2",
            },
        )
        self.assertEqual(saved.status_code, 302)
        faculty_client = app_module.app.test_client()
        self.set_role(faculty_client, "FAC100", "faculty")
        schedule = faculty_client.get("/faculty/schedule")
        self.assertEqual(schedule.status_code, 200)
        self.assertIn(b"Data Structures", schedule.data)
        self.assertIn(b"C-204", schedule.data)
        student_client = app_module.app.test_client()
        self.set_role(student_client, "STU001", "student")
        student_schedule = student_client.get("/student/schedule")
        self.assertIn(b"Data Structures", student_schedule.data)
        other_class_client = app_module.app.test_client()
        self.set_role(other_class_client, "STU003", "student")
        self.assertNotIn(b"Data Structures", other_class_client.get("/student/schedule").data)

    def test_hod_assigns_multiple_classes_and_faculty_marks_selected_roster(self):
        self.add_approved_faculty()
        self.set_role(self.client, "cshod", "hod")
        assigned = self.client.post(
            "/hod/classes",
            data={
                "action": "assign-class",
                "faculty_id": "FAC100",
                "branch": "CSE",
                "year": "3",
                "section": "B",
            },
        )
        self.assertEqual(assigned.status_code, 302)
        hod_classes_page = self.client.get("/hod/classes")
        self.assertEqual(hod_classes_page.status_code, 200)
        self.assertIn(b"Faculty class assignments", hod_classes_page.data)
        connection = sqlite3.connect(self.database_path)
        class_id = connection.execute(
            "SELECT id FROM faculty_classes WHERE faculty_id='FAC100' AND year='3' AND section='B'"
        ).fetchone()[0]
        connection.close()
        scheduled = self.client.post(
            "/hod/classes",
            data={
                "class_assignment_id": str(class_id),
                "subject": "Algorithms",
                "room": "C-305",
                "weekday": "1",
                "period": "3",
            },
        )
        self.assertEqual(scheduled.status_code, 302)
        original_db_path = database_module.DB_PATH
        database_module.DB_PATH = self.database_path
        try:
            room_roster = database_module.get_scheduled_roster("C-305", "2026-09-29", 3)
        finally:
            database_module.DB_PATH = original_db_path
        self.assertEqual(room_roster, ["STU005"])

        self.set_role(self.client, "FAC100", "faculty")
        selected_roster = self.client.get(f"/faculty?class_id={class_id}")
        self.assertIn(b"Student Five", selected_roster.data)
        self.assertNotIn(b"Student One", selected_roster.data)
        student_details = self.client.get(f"/faculty/students?class_id={class_id}")
        self.assertIn(b"Student Five", student_details.data)
        self.assertNotIn(b"Student One", student_details.data)
        schedule = self.client.get(f"/faculty/schedule?class_id={class_id}")
        self.assertIn(b"Algorithms", schedule.data)
        self.assertIn(b"C-305", schedule.data)
        coursework = self.client.get(f"/faculty/coursework?class_id={class_id}")
        self.assertIn(b"Year 3", coursework.data)
        assignment_response = self.client.post(
            "/faculty/coursework",
            data={
                "class_id": str(class_id),
                "action": "assignment",
                "title": "Algorithms exercise",
                "description": "Complete the worksheet.",
                "due_at": "2026-10-01T09:00",
            },
        )
        self.assertEqual(assignment_response.status_code, 302)
        connection = sqlite3.connect(self.database_path)
        assignment_class = connection.execute(
            "SELECT branch, year, section FROM assignments WHERE title='Algorithms exercise'"
        ).fetchone()
        connection.close()
        self.assertEqual(assignment_class, ("CSE", "3", "B"))
        saved = self.client.post(
            "/faculty",
            data={
                "class_id": str(class_id),
                "date": "2026-09-29",
                "period": "2",
                "attendance_STU005": "present",
                "attendance_STU001": "present",
            },
        )
        self.assertEqual(saved.status_code, 302)
        connection = sqlite3.connect(self.database_path)
        records = connection.execute(
            "SELECT student_id, status FROM period_attendance WHERE date='2026-09-29' AND period=2 ORDER BY student_id"
        ).fetchall()
        connection.close()
        self.assertEqual(records, [("STU005", "present")])

    def test_assignment_upload_is_visible_to_class_and_download_is_scoped(self):
        from io import BytesIO

        self.add_approved_faculty()
        self.set_role(self.client, "FAC100", "faculty")
        rejected_announcement = self.client.post(
            "/faculty/coursework",
            data={"action": "announcement", "title": "College update", "body": "Read the notice."},
        )
        self.assertEqual(rejected_announcement.status_code, 302)
        hod_client = app_module.app.test_client()
        self.set_role(hod_client, "cshod", "hod")
        hod_announcement = hod_client.post(
            "/hod/announcements",
            data={"title": "College update", "body": "Read the notice."},
        )
        self.assertEqual(hod_announcement.status_code, 302)
        self.client.post(
            "/faculty/coursework",
            data={
                "action": "assignment",
                "title": "Week one",
                "description": "Submit the exercise.",
                "due_at": "2026-10-01T09:00",
            },
        )
        connection = sqlite3.connect(self.database_path)
        assignment_id = connection.execute("SELECT id FROM assignments").fetchone()[0]
        connection.close()

        student_client = app_module.app.test_client()
        self.set_role(student_client, "STU001", "student")
        page = student_client.get("/student/coursework")
        self.assertIn(b"College update", page.data)
        self.assertIn(b"Week one", page.data)
        self.assertIn(b"College update", hod_client.get("/hod/announcements").data)
        self.assertIn(b"College update", self.client.get("/faculty/coursework").data)
        with patch.object(app_module.app, "root_path", self.temp_dir.name):
            uploaded = student_client.post(
                f"/student/coursework/{assignment_id}/submit",
                data={"submission": (BytesIO(b"student work"), "week-one.pdf")},
                content_type="multipart/form-data",
            )
            self.assertEqual(uploaded.status_code, 302)
            connection = sqlite3.connect(self.database_path)
            stored_filename = connection.execute(
                "SELECT stored_filename FROM assignment_submissions WHERE assignment_id=? AND student_id='STU001'",
                (assignment_id,),
            ).fetchone()[0]
            connection.close()
            download = student_client.get(f"/coursework-files/{stored_filename}")
            self.assertEqual(download.status_code, 200)
            self.assertEqual(download.data, b"student work")

            other_student = app_module.app.test_client()
            self.set_role(other_student, "STU003", "student")
            other_class_page = other_student.get("/student/coursework")
            self.assertIn(b"College update", other_class_page.data)
            self.assertNotIn(b"Week one", other_class_page.data)
            self.assertEqual(other_student.get(f"/coursework-files/{stored_filename}").status_code, 404)

            faculty_client = app_module.app.test_client()
            self.set_role(faculty_client, "FAC100", "faculty")
            faculty_download = faculty_client.get(f"/coursework-files/{stored_filename}")
            self.assertEqual(faculty_download.status_code, 200)

    def test_hod_announcements_are_filtered_by_recipient(self):
        self.add_approved_faculty()
        self.set_role(self.client, "cshod", "hod")
        for title, audience in (("For students", "students"), ("For faculty", "faculty"), ("For all", "all")):
            response = self.client.post(
                "/hod/announcements",
                data={"title": title, "body": f"{title} notice", "audience": audience},
            )
            self.assertEqual(response.status_code, 302)

        hod_page = self.client.get("/hod/announcements")
        self.assertIn(b"For students", hod_page.data)
        self.assertIn(b"For faculty", hod_page.data)
        self.assertIn(b"For all", hod_page.data)

        faculty_client = app_module.app.test_client()
        self.set_role(faculty_client, "FAC100", "faculty")
        faculty_page = faculty_client.get("/faculty/coursework")
        self.assertNotIn(b"For students", faculty_page.data)
        self.assertIn(b"For faculty", faculty_page.data)
        self.assertIn(b"For all", faculty_page.data)

        student_client = app_module.app.test_client()
        self.set_role(student_client, "STU001", "student")
        student_page = student_client.get("/student/coursework")
        self.assertIn(b"For students", student_page.data)
        self.assertNotIn(b"For faculty", student_page.data)
        self.assertIn(b"For all", student_page.data)

    def test_excel_export_separates_department_students_and_faculty(self):
        connection = sqlite3.connect(self.database_path)
        connection.executemany(
            "INSERT INTO users (user_id, password, role, dob, branch, name, year, section, faculty_approved, email, phone) "
            "VALUES (?, ?, 'faculty', NULL, ?, ?, ?, ?, 1, NULL, NULL)",
            [
                ("FAC-CSE", "pw", "CSE", "Faculty CSE", "2", "A"),
                ("FAC-ECE", "pw", "ECE", "Faculty ECE", "2", "A"),
            ],
        )
        connection.executemany(
            "INSERT INTO period_attendance (student_id, date, period, status, value, first_seen_time) "
            "VALUES (?, '2026-09-29', 3, 'present', 1.0, '10:00:00')",
            [("STU001",), ("STU004",)],
        )
        connection.commit()
        connection.close()
        self.set_role(self.client, "cshod", "hod")

        response = self.client.get("/hod/export-excel?filter=all")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Students", "Faculty"])
        student_rows = list(workbook["Students"].values)
        faculty_rows = list(workbook["Faculty"].values)
        self.assertTrue(any(row[0] == "STU001" for row in student_rows))
        self.assertFalse(any(row[0] == "STU004" for row in student_rows))
        self.assertTrue(any(row[0] == "FAC-CSE" for row in faculty_rows))
        self.assertFalse(any(row[0] == "FAC-ECE" for row in faculty_rows))
        workbook.close()


if __name__ == "__main__":
    unittest.main()